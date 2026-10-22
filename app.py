"""TRỌ VIỆT MANAGER — Streamlit owner-operated rental billing.

Based on the original 2026 Streamlit app.py (month/year selection,
old/new meter readings, rent + utility computation, CSV export).
This version adds durable local SQLite ledger, invoice snapshots,
payments, PDF notices/receipts, Excel/PDF reports and optional SMTP.

SECURITY: This is a single-owner app. Set APP_PASSWORD in Streamlit
secrets before putting it on a public URL. SQLite on ephemeral hosts is
NOT durable cloud storage; see README.md and in-app Cài đặt section.
"""
from __future__ import annotations

import calendar
import csv
import hashlib
import hmac
import io
import json
import os
import re
import smtplib
import sqlite3
import ssl
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from email.message import EmailMessage
from html import escape
from pathlib import Path
from typing import Any
from uuid import uuid4
from zoneinfo import ZoneInfo

APP_VERSION = "6.1.1-local"
DB_PATH = os.environ.get("TRO_VIET_DB", "tro_viet.sqlite3")
CATEGORIES = {"Tiêu chuẩn": 2_000_000, "Tiện nghi": 3_200_000,
              "Cao cấp": 4_500_000, "Full nội thất": 8_000_000}
MENUS = ["Tổng quan", "Phòng trọ", "Điện & nước", "Thu tiền", "Báo cáo", "Cài đặt"]
MONEY_FIELDS = {"rent_amount", "electric_amount", "water_amount", "fee_amount",
                "discount_amount", "amount", "paid", "outstanding"}
BACKUP_TABLES = ["settings", "rooms", "contracts", "readings", "invoices",
                 "payments", "allocations", "payment_reversals", "email_events", "audit_logs"]


def vn_today() -> date:
    return datetime.now(ZoneInfo("Asia/Ho_Chi_Minh")).date()


def now_iso() -> str:
    return datetime.now(ZoneInfo("Asia/Ho_Chi_Minh")).isoformat(timespec="seconds")


def money(v: Any) -> str:
    return f"{int(v or 0):,}".replace(",", ".") + "đ"


def dec(value: Any) -> Decimal:
    try:
        val = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        raise ValueError("Giá trị công tơ không hợp lệ")
    if not val.is_finite() or val < 0:
        raise ValueError("Chỉ số hoặc số lượng không được âm")
    return val


def to_vnd(value: Decimal | int | float) -> int:
    return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def next_period(period: str) -> str:
    y, m = map(int, period.split("-"))
    return f"{y + (m == 12):04d}-{m % 12 + 1:02d}"


def previous_period(period: str) -> str:
    y, m = map(int, period.split("-"))
    return f"{y - (m == 1):04d}-{12 if m == 1 else m - 1:02d}"


def period_dates(period: str) -> tuple[date, date]:
    y, m = map(int, period.split("-"))
    return date(y, m, 1), date(y, m, calendar.monthrange(y, m)[1])


def period_label(period: str) -> str:
    y, m = period.split("-")
    return f"Tháng {int(m):02d}/{y}"


def due_date(period: str, day: int = 5) -> str:
    nxt = next_period(period)
    y, m = map(int, nxt.split("-"))
    return date(y, m, min(day, calendar.monthrange(y, m)[1])).isoformat()


def email_valid(address: str) -> bool:
    return bool(re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", (address or "").strip()))


def normalize_date(value: Any, field: str = "Ngày") -> str:
    """Accept Vietnamese dd/mm/yyyy and ISO yyyy-mm-dd; persist ISO only.

    Empty optional dates remain empty.  Raises a human-readable validation
    error rather than leaking date.fromisoformat's ValueError to the UI.
    """
    if value is None or str(value).strip() == "":
        return ""
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    raw = str(value).strip()
    try:
        if "/" in raw:
            parsed = datetime.strptime(raw, "%d/%m/%Y").date()
        else:
            parsed = date.fromisoformat(raw)
        return parsed.isoformat()
    except (ValueError, TypeError):
        raise ValueError(f"{field} không hợp lệ: '{raw}'. Hãy nhập theo ngày/tháng/năm, ví dụ 15/01/2026.") from None


def display_date(value: Any) -> str:
    """Show dates as dd/mm/yyyy while retaining ISO strings in SQLite."""
    normalized = normalize_date(value)
    return date.fromisoformat(normalized).strftime("%d/%m/%Y") if normalized else ""


def rent_for_period(monthly: int, period: str, starts: str, ends: str) -> int:
    first, last = period_dates(period)
    effective_start = max(first, date.fromisoformat(normalize_date(starts, "Ngày bắt đầu"))) if starts else first
    effective_end = min(last, date.fromisoformat(normalize_date(ends, "Ngày kết thúc"))) if ends else last
    if effective_end < effective_start:
        return 0
    active_days = (effective_end - effective_start).days + 1
    return to_vnd(Decimal(monthly) * Decimal(active_days) / Decimal(last.day))


def calculate_bill(rent: int, e_start: Any, e_end: Any, w_start: Any,
                   w_end: Any, e_rate: int, w_rate: int, fee: int,
                   discount: int = 0) -> dict[str, Any]:
    e_use, w_use = dec(e_end) - dec(e_start), dec(w_end) - dec(w_start)
    if e_use < 0 or w_use < 0:
        raise ValueError("Chỉ số mới thấp hơn chỉ số cũ; hãy kiểm tra/thay công tơ trước")
    electric = to_vnd(e_use * Decimal(e_rate))
    water = to_vnd(w_use * Decimal(w_rate))
    total = int(rent) + electric + water + int(fee) - int(discount)
    if total < 0:
        raise ValueError("Giảm trừ không được vượt quá tổng tiền phát sinh")
    return {"electric_use": str(e_use), "water_use": str(w_use),
            "electric_amount": electric, "water_amount": water,
            "rent_amount": int(rent), "fee_amount": int(fee),
            "discount_amount": int(discount), "amount": total}


@contextmanager
def db_transaction(db_path: str = DB_PATH):
    path = Path(db_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(path), timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    conn.execute("PRAGMA busy_timeout=30000")
    try:
        conn.execute("BEGIN IMMEDIATE")
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def query(sql: str, params: tuple = (), db_path: str = DB_PATH) -> list[dict[str, Any]]:
    with db_transaction(db_path) as con:
        return [dict(x) for x in con.execute(sql, params).fetchall()]


def init_db(db_path: str = DB_PATH):
    schema = """
    CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS rooms (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        code TEXT NOT NULL UNIQUE, category TEXT NOT NULL DEFAULT 'Tiêu chuẩn',
        rent INTEGER NOT NULL DEFAULT 2000000 CHECK(rent >= 0),
        fee INTEGER NOT NULL DEFAULT 0 CHECK(fee >= 0),
        discount INTEGER NOT NULL DEFAULT 0 CHECK(discount >= 0),
        status TEXT NOT NULL DEFAULT 'Còn trống',
        tenant TEXT NOT NULL DEFAULT '', email TEXT NOT NULL DEFAULT '',
        phone TEXT NOT NULL DEFAULT '', start_date TEXT NOT NULL DEFAULT '',
        end_date TEXT NOT NULL DEFAULT '', deposit INTEGER NOT NULL DEFAULT 0,
        building TEXT NOT NULL DEFAULT '', floor TEXT NOT NULL DEFAULT '',
        area TEXT NOT NULL DEFAULT '', notes TEXT NOT NULL DEFAULT '',
        initial_e TEXT NOT NULL DEFAULT '', initial_w TEXT NOT NULL DEFAULT '',
        deleted_at TEXT
    );
    CREATE TABLE IF NOT EXISTS contracts (
        id INTEGER PRIMARY KEY AUTOINCREMENT, room_id INTEGER NOT NULL REFERENCES rooms(id),
        tenant TEXT NOT NULL, email TEXT NOT NULL, start_date TEXT NOT NULL,
        end_date TEXT NOT NULL, monthly_rent INTEGER NOT NULL, deposit INTEGER NOT NULL,
        saved_at TEXT NOT NULL, note TEXT NOT NULL DEFAULT ''
    );
    CREATE TABLE IF NOT EXISTS readings (
        room_id INTEGER NOT NULL REFERENCES rooms(id), period TEXT NOT NULL,
        e_start TEXT NOT NULL, e_end TEXT NOT NULL,
        w_start TEXT NOT NULL, w_end TEXT NOT NULL,
        updated_at TEXT NOT NULL, PRIMARY KEY(room_id,period)
    );
    CREATE TABLE IF NOT EXISTS invoices (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        room_id INTEGER NOT NULL REFERENCES rooms(id), period TEXT NOT NULL,
        issued_at TEXT NOT NULL, due_at TEXT NOT NULL,
        snapshot TEXT NOT NULL, amount INTEGER NOT NULL CHECK(amount>=0),
        prior_balance INTEGER NOT NULL DEFAULT 0, credit_used INTEGER NOT NULL DEFAULT 0,
        UNIQUE(room_id,period)
    );
    CREATE TABLE IF NOT EXISTS payments (
        id INTEGER PRIMARY KEY AUTOINCREMENT, room_id INTEGER NOT NULL REFERENCES rooms(id),
        amount INTEGER NOT NULL CHECK(amount>0), paid_at TEXT NOT NULL,
        method TEXT NOT NULL, note TEXT NOT NULL DEFAULT '',
        actor TEXT NOT NULL DEFAULT 'Chủ trọ', token TEXT NOT NULL UNIQUE
    );
    CREATE TABLE IF NOT EXISTS allocations (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        payment_id INTEGER NOT NULL REFERENCES payments(id),
        invoice_id INTEGER NOT NULL REFERENCES invoices(id),
        amount INTEGER NOT NULL CHECK(amount>0), UNIQUE(payment_id,invoice_id)
    );
    CREATE TABLE IF NOT EXISTS payment_reversals (
        payment_id INTEGER PRIMARY KEY REFERENCES payments(id),
        reversed_at TEXT NOT NULL, reason TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS email_events (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        invoice_id INTEGER NOT NULL REFERENCES invoices(id),
        recipient TEXT NOT NULL, status TEXT NOT NULL,
        message TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS audit_logs (
        id INTEGER PRIMARY KEY AUTOINCREMENT, event TEXT NOT NULL,
        details TEXT NOT NULL, created_at TEXT NOT NULL
    );
    CREATE INDEX IF NOT EXISTS idx_invoice_period ON invoices(period);
    CREATE INDEX IF NOT EXISTS idx_payments_room ON payments(room_id);
    CREATE INDEX IF NOT EXISTS idx_readings_period ON readings(period);
    """
    with db_transaction(db_path) as c:
        c.executescript(schema)
        defaults = {"property_name": "Nhà trọ của tôi", "property_address": "",
                    "property_phone": "", "electric_rate": "3500", "water_rate": "20000",
                    "bank_name": "", "bank_account": "", "bank_holder": "",
                    "due_day": "5", "auto_email": "0"}
        for k, v in defaults.items():
            c.execute("INSERT OR IGNORE INTO settings(key,value) VALUES(?,?)", (k,v))
        count = c.execute("SELECT COUNT(*) FROM rooms").fetchone()[0]
        if count == 0:
            c.executemany("INSERT INTO rooms(code) VALUES(?)", [(f"P{i:03d}",) for i in range(1,57)])


def audit(c, event: str, details: Any):
    c.execute("INSERT INTO audit_logs(event,details,created_at) VALUES(?,?,?)",
              (event, json.dumps(details, ensure_ascii=False, default=str), now_iso()))


def settings_get(db_path: str = DB_PATH) -> dict[str, str]:
    return {r["key"]: r["value"] for r in query("SELECT key,value FROM settings", db_path=db_path)}


def settings_save(values: dict[str, str], db_path: str = DB_PATH):
    valid = {"property_name", "property_address", "property_phone", "electric_rate",
             "water_rate", "bank_name", "bank_account", "bank_holder", "due_day", "auto_email"}
    with db_transaction(db_path) as c:
        for k,v in values.items():
            if k not in valid:
                raise ValueError("Tùy chọn không hợp lệ")
            c.execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(k,str(v)))
        audit(c, "save_settings", {"keys":list(values)})


def rooms_list(db_path: str = DB_PATH, include_deleted: bool=False):
    return query("SELECT * FROM rooms" + ("" if include_deleted else " WHERE deleted_at IS NULL") + " ORDER BY code", db_path=db_path)


def save_room(data: dict[str, Any], db_path: str = DB_PATH) -> int:
    # Normalize before comparisons, contract snapshots and database writes.
    # Keep the caller's dictionary untouched.
    data = dict(data)
    data["start_date"] = normalize_date(data.get("start_date"), "Ngày bắt đầu")
    data["end_date"] = normalize_date(data.get("end_date"), "Ngày kết thúc")
    required = ["code", "category", "rent", "fee", "discount", "status", "tenant", "email",
                "phone", "start_date", "end_date", "deposit", "building", "floor", "area",
                "notes", "initial_e", "initial_w"]
    if not str(data.get("code", "")).strip():
        raise ValueError("Cần nhập mã phòng")
    if data.get("email") and not email_valid(str(data["email"])):
        raise ValueError("Email khách thuê không hợp lệ")
    for k in ("rent","fee","discount","deposit"):
        if int(data.get(k,0)) < 0:
            raise ValueError(f"{k} không được âm")
    if data.get("start_date") and data.get("end_date") and data["start_date"] > data["end_date"]:
        raise ValueError("Ngày kết thúc phải sau ngày bắt đầu")
    for k in ("initial_e", "initial_w"):
        if str(data.get(k, "")).strip(): dec(data[k])
    cols = [str(data.get(k,"")) if k not in ("rent","fee","discount","deposit") else int(data.get(k,0)) for k in required]
    with db_transaction(db_path) as c:
        room_id = data.get("id")
        old = c.execute("SELECT * FROM rooms WHERE id=?",(room_id,)).fetchone() if room_id else None
        if old:
            c.execute("UPDATE rooms SET " + ",".join(f"{k}=?" for k in required) + " WHERE id=?",(*cols,room_id))
        else:
            cur = c.execute("INSERT INTO rooms(" + ",".join(required) + ") VALUES(" + ",".join("?" for _ in required) + ")",cols)
            room_id = cur.lastrowid
        snapshot_fields = ("tenant", "email", "start_date", "end_date", "rent", "deposit")
        if data.get("tenant") and (old is None or any(str(old[k]) != str(data.get(k, "")) for k in snapshot_fields)):
            c.execute("INSERT INTO contracts(room_id,tenant,email,start_date,end_date,monthly_rent,deposit,saved_at,note) VALUES(?,?,?,?,?,?,?,?,?)",
                      (room_id,data["tenant"],data["email"],data["start_date"],data["end_date"],int(data["rent"]),int(data["deposit"]),now_iso(),"Lưu phiên bản thông tin thuê"))
        audit(c, "save_room", {"room_id":room_id,"code":data["code"]})
        return int(room_id)


def delete_room(room_id: int, db_path: str = DB_PATH):
    with db_transaction(db_path) as c:
        c.execute("UPDATE rooms SET deleted_at=? WHERE id=?",(now_iso(),room_id))
        audit(c,"soft_delete_room",{"room_id":room_id})


def _latest_meter(c, room_id: int, period: str, kind: str) -> str:
    row = c.execute(f"SELECT json_extract(snapshot, '$.{kind}_end') AS last_value FROM invoices WHERE room_id=? AND period<? ORDER BY period DESC LIMIT 1",(room_id,period)).fetchone()
    if row and row["last_value"] is not None:
        return str(row["last_value"])
    row = c.execute("SELECT initial_e, initial_w FROM rooms WHERE id=?",(room_id,)).fetchone()
    return str(row["initial_e" if kind == "e" else "initial_w"]) if row else ""


def meter_start(room_id: int, period: str, kind: str, db_path: str = DB_PATH) -> str:
    with db_transaction(db_path) as c:
        current = c.execute("SELECT e_start,w_start FROM readings WHERE room_id=? AND period=?",(room_id,period)).fetchone()
        if current:
            return current["e_start" if kind == "e" else "w_start"]
        return _latest_meter(c,room_id,period,kind)


def save_reading(room_id: int, period: str, e_start: Any, e_end: Any,
                 w_start: Any, w_end: Any, db_path: str = DB_PATH):
    e0,e1,w0,w1 = (dec(x) for x in (e_start,e_end,w_start,w_end))
    if e1 < e0 or w1 < w0:
        raise ValueError("Số mới nhỏ hơn số cũ. Cần xử lý thay/đảo công tơ riêng.")
    with db_transaction(db_path) as c:
        issued = c.execute("SELECT 1 FROM invoices WHERE room_id=? AND period=?",(room_id,period)).fetchone()
        later = c.execute("SELECT 1 FROM invoices WHERE room_id=? AND period>? LIMIT 1",(room_id,period)).fetchone()
        if issued or later:
            raise ValueError("Kỳ này hoặc kỳ sau đã chốt chứng từ: không thể sửa âm thầm chỉ số")
        # Preserve the start reading chosen by the manager; validate continuity with earlier invoices.
        previous_e = _latest_meter(c,room_id,period,"e")
        previous_w = _latest_meter(c,room_id,period,"w")
        if (previous_e != "" and e0 != dec(previous_e)) or (previous_w != "" and w0 != dec(previous_w)):
            raise ValueError("Số đầu kỳ không khớp kỳ đã chốt/cấu hình. Hãy kiểm tra trước khi lưu.")
        c.execute("INSERT INTO readings(room_id,period,e_start,e_end,w_start,w_end,updated_at) VALUES(?,?,?,?,?,?,?) ON CONFLICT(room_id,period) DO UPDATE SET e_start=excluded.e_start,e_end=excluded.e_end,w_start=excluded.w_start,w_end=excluded.w_end,updated_at=excluded.updated_at",
                  (room_id,period,str(e0),str(e1),str(w0),str(w1),now_iso()))
        audit(c,"meter_reading",{"room_id":room_id,"period":period})


def _alloc_rows(c, room_id: int) -> list[sqlite3.Row]:
    return c.execute("""SELECT i.id, i.amount,
        COALESCE((SELECT SUM(a.amount) FROM allocations a JOIN payments p ON p.id=a.payment_id
        LEFT JOIN payment_reversals r ON r.payment_id=p.id
        WHERE a.invoice_id=i.id AND r.payment_id IS NULL),0) AS paid
        FROM invoices i WHERE i.room_id=? ORDER BY i.period,i.id""",(room_id,)).fetchall()


def _allocate_credits(c, room_id: int):
    invoices = _alloc_rows(c,room_id)
    credits = c.execute("""SELECT p.id,p.amount,COALESCE(SUM(a.amount),0) AS allocated
        FROM payments p LEFT JOIN allocations a ON a.payment_id=p.id
        LEFT JOIN payment_reversals r ON r.payment_id=p.id
        WHERE p.room_id=? AND r.payment_id IS NULL
        GROUP BY p.id ORDER BY p.paid_at,p.id""",(room_id,)).fetchall()
    inv_rest = [[r["id"],max(0,int(r["amount"])-int(r["paid"]))] for r in invoices]
    for pay in credits:
        remaining = int(pay["amount"]) - int(pay["allocated"])
        for inv in inv_rest:
            used = min(remaining,inv[1])
            if used:
                c.execute("""INSERT INTO allocations(payment_id,invoice_id,amount) VALUES(?,?,?)
                    ON CONFLICT(payment_id,invoice_id) DO UPDATE SET amount=amount+excluded.amount""",
                          (pay["id"],inv[0],used))
                remaining -= used; inv[1] -= used
            if remaining <= 0: break


def debt_room(room_id: int, db_path: str = DB_PATH) -> int:
    with db_transaction(db_path) as c:
        return sum(max(0, int(r["amount"])-int(r["paid"])) for r in _alloc_rows(c,room_id))


def credit_room(room_id: int, db_path: str = DB_PATH) -> int:
    with db_transaction(db_path) as c:
        row=c.execute("""SELECT COALESCE(SUM(p.amount),0) AS amount,
         COALESCE(SUM((SELECT COALESCE(SUM(a.amount),0) FROM allocations a WHERE a.payment_id=p.id)),0) AS allocated
         FROM payments p LEFT JOIN payment_reversals r ON r.payment_id=p.id
         WHERE p.room_id=? AND r.payment_id IS NULL""",(room_id,)).fetchone()
        return int(row["amount"])-int(row["allocated"])


def issue_invoice(room_id: int, period: str, db_path: str = DB_PATH) -> int:
    with db_transaction(db_path) as c:
        exists=c.execute("SELECT id FROM invoices WHERE room_id=? AND period=?",(room_id,period)).fetchone()
        if exists:
            raise ValueError("Phòng này đã chốt kỳ. Không tạo lại thông báo trùng.")
        room=c.execute("SELECT * FROM rooms WHERE id=? AND deleted_at IS NULL",(room_id,)).fetchone()
        if not room or room["status"] != "Đang thuê" or not room["tenant"]:
            raise ValueError("Cần có phòng đang thuê và tên khách trước khi chốt")
        reading=c.execute("SELECT * FROM readings WHERE room_id=? AND period=?",(room_id,period)).fetchone()
        if not reading:
            raise ValueError("Phòng chưa nhập đủ chỉ số điện nước kỳ này")
        cfg={r["key"]:r["value"] for r in c.execute("SELECT key,value FROM settings")}
        rent=rent_for_period(int(room["rent"]),period,room["start_date"],room["end_date"])
        if rent == 0:
            raise ValueError("Hợp đồng không nằm trong kỳ đang chốt")
        detail=calculate_bill(rent,reading["e_start"],reading["e_end"],reading["w_start"],reading["w_end"],
                              int(cfg["electric_rate"]),int(cfg["water_rate"]),int(room["fee"]),int(room["discount"]))
        prior=sum(max(0,int(r["amount"])-int(r["paid"])) for r in _alloc_rows(c,room_id))
        prev_credit=c.execute("""SELECT COALESCE(SUM(p.amount),0)-COALESCE(SUM((SELECT COALESCE(SUM(a.amount),0) FROM allocations a WHERE a.payment_id=p.id)),0)
            FROM payments p LEFT JOIN payment_reversals r ON r.payment_id=p.id
            WHERE p.room_id=? AND r.payment_id IS NULL""",(room_id,)).fetchone()[0]
        snapshot={"room_code":room["code"],"tenant":room["tenant"],"email":room["email"],
                  "property_name":cfg["property_name"],"property_address":cfg["property_address"],
                  "property_phone":cfg["property_phone"],"bank_name":cfg["bank_name"],
                  "bank_account":cfg["bank_account"],"bank_holder":cfg["bank_holder"],
                  "period":period,"e_start":reading["e_start"],"e_end":reading["e_end"],
                  "w_start":reading["w_start"],"w_end":reading["w_end"],
                  "e_rate":int(cfg["electric_rate"]),"w_rate":int(cfg["water_rate"]),
                  "contract_start":room["start_date"],"contract_end":room["end_date"],
                  **detail}
        cur=c.execute("INSERT INTO invoices(room_id,period,issued_at,due_at,snapshot,amount,prior_balance,credit_used) VALUES(?,?,?,?,?,?,?,?)",
                      (room_id,period,now_iso(),due_date(period,int(cfg["due_day"])),
                       json.dumps(snapshot,ensure_ascii=False),detail["amount"],prior,min(int(prev_credit),detail["amount"])))
        _allocate_credits(c,room_id)
        audit(c,"issue_invoice",{"invoice_id":cur.lastrowid,"room_id":room_id,"period":period})
        return int(cur.lastrowid)


def invoice_rows(period: str|None=None, db_path: str = DB_PATH):
    sql="""SELECT i.*,r.code,r.tenant AS current_tenant,
    COALESCE((SELECT SUM(a.amount) FROM allocations a JOIN payments p ON p.id=a.payment_id
     LEFT JOIN payment_reversals v ON v.payment_id=p.id
     WHERE a.invoice_id=i.id AND v.payment_id IS NULL),0) AS paid
    FROM invoices i JOIN rooms r ON r.id=i.room_id"""
    if period:
        sql+=" WHERE i.period=?"
    sql+=" ORDER BY i.period DESC,r.code"
    rows=query(sql,(period,) if period else (),db_path)
    for r in rows:
        r["snapshot"]=json.loads(r["snapshot"])
        r["outstanding"]=max(0,int(r["amount"])-int(r["paid"]))
    return rows


def invoice_by_id(invoice_id: int, db_path: str=DB_PATH) -> dict[str, Any]:
    with db_transaction(db_path) as c:
        row=c.execute("SELECT * FROM invoices WHERE id=?",(invoice_id,)).fetchone()
        if not row: raise ValueError("Không tìm thấy thông báo")
        value=dict(row);value["snapshot"]=json.loads(value["snapshot"])
        value["paid"]=sum(int(r["amount"]) for r in c.execute("""SELECT a.amount FROM allocations a
           JOIN payments p ON p.id=a.payment_id LEFT JOIN payment_reversals v ON v.payment_id=p.id
           WHERE a.invoice_id=? AND v.payment_id IS NULL""",(invoice_id,)))
        return value


def record_payment(room_id: int, amount: int, paid_at: str, method: str, note: str="",
                   token: str="", db_path: str=DB_PATH) -> int:
    if not int(amount)>0: raise ValueError("Số tiền phải lớn hơn 0")
    paid_at = normalize_date(paid_at, "Ngày thanh toán")
    if not paid_at:
        raise ValueError("Cần nhập ngày thanh toán")
    token=token or uuid4().hex
    with db_transaction(db_path) as c:
        row=c.execute("SELECT id FROM payments WHERE token=?",(token,)).fetchone()
        if row: return int(row["id"])
        room=c.execute("SELECT id FROM rooms WHERE id=? AND deleted_at IS NULL",(room_id,)).fetchone()
        if not room: raise ValueError("Phòng không tồn tại")
        cur=c.execute("INSERT INTO payments(room_id,amount,paid_at,method,note,token) VALUES(?,?,?,?,?,?)",
                      (room_id,int(amount),paid_at,method,note,token))
        _allocate_credits(c,room_id)
        audit(c,"record_payment",{"payment_id":cur.lastrowid,"room_id":room_id,"amount":int(amount)})
        return int(cur.lastrowid)


def reverse_payment(payment_id: int, reason: str, db_path: str=DB_PATH):
    if len(reason.strip())<5: raise ValueError("Cần nêu lý do điều chỉnh (ít nhất 5 ký tự)")
    with db_transaction(db_path) as c:
        row=c.execute("SELECT room_id FROM payments WHERE id=?",(payment_id,)).fetchone()
        if not row: raise ValueError("Không tìm thấy giao dịch")
        if c.execute("SELECT 1 FROM payment_reversals WHERE payment_id=?",(payment_id,)).fetchone():
            raise ValueError("Giao dịch này đã được đảo trước đó")
        c.execute("INSERT INTO payment_reversals(payment_id,reversed_at,reason) VALUES(?,?,?)",(payment_id,now_iso(),reason.strip()))
        audit(c,"reverse_payment",{"payment_id":payment_id,"reason":reason.strip()})
        _allocate_credits(c,row["room_id"])


def payments_list(db_path: str=DB_PATH):
    return query("""SELECT p.*,r.code,v.reversed_at,v.reason AS reversal_reason
                   FROM payments p JOIN rooms r ON r.id=p.room_id
                   LEFT JOIN payment_reversals v ON v.payment_id=p.id ORDER BY p.id DESC""",db_path=db_path)


def report_metrics(period: str, db_path: str=DB_PATH):
    rows=invoice_rows(period,db_path)
    rooms=rooms_list(db_path)
    settled=sum(1 for i in rows if i["outstanding"]==0)
    part=sum(1 for i in rows if 0<i["paid"]<i["amount"])
    outstanding_all=sum(r["outstanding"] for r in invoice_rows(db_path=db_path))
    cash=int(query("""SELECT COALESCE(SUM(p.amount),0) AS v FROM payments p
        LEFT JOIN payment_reversals v ON v.payment_id=p.id
        WHERE v.payment_id IS NULL AND substr(p.paid_at,1,7)=?""",(period,),db_path)[0]["v"])
    return {"rooms":len(rooms),"occupied":sum(r["status"]=="Đang thuê" for r in rooms),
            "invoices":len(rows),"issued":sum(int(i["amount"]) for i in rows),
            "paid_on_invoices":sum(int(i["paid"]) for i in rows),
            "cash_in_period":cash,"outstanding_all":outstanding_all,
            "paid_full":settled,"paid_part":part,"unpaid":len(rows)-settled-part}

# --------------------------------------------------------------
# Real PDF / XLSX / CSV exports
# --------------------------------------------------------------

def _pdf_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    paths = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/local/share/fonts/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
        Path("assets/DejaVuSans.ttf"),
    ]
    for path in paths:
        if path.exists():
            if "VietSans" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("VietSans",str(path)))
            return "VietSans"
    raise RuntimeError("Không tìm thấy font Unicode để xuất PDF tiếng Việt. Cài DejaVu Sans trên máy chủ.")


def _pdf_doc(title: str, subtitle: str, sections: list[tuple[str,list[list[str]]]],
             footer: str="", landscape_mode: bool=False) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, KeepTogether
    font=_pdf_font()
    b=io.BytesIO()
    page=landscape(A4) if landscape_mode else A4
    doc=SimpleDocTemplate(b,pagesize=page,leftMargin=15*mm,rightMargin=15*mm,
                          topMargin=15*mm,bottomMargin=16*mm)
    style=ParagraphStyle("VietTitle",fontName=font,fontSize=15,leading=21,
                         alignment=TA_CENTER,textColor=colors.HexColor("#17324D"),spaceAfter=7)
    sub=ParagraphStyle("VietSub",fontName=font,fontSize=9.2,leading=14,
                       alignment=TA_CENTER,textColor=colors.HexColor("#475569"))
    cell=ParagraphStyle("VietCell",fontName=font,fontSize=8,leading=11)
    caption=ParagraphStyle("VietHeading",fontName=font,fontSize=10,leading=15,
                           textColor=colors.HexColor("#1769C2"),spaceBefore=10,spaceAfter=5)
    foot=ParagraphStyle("VietFoot",fontName=font,fontSize=8,leading=11,
                        textColor=colors.HexColor("#475569"),spaceBefore=12)
    esc=lambda s: escape(str(s)).replace("\n","<br/>")
    story=[Paragraph(esc(title),style),Paragraph(esc(subtitle),sub),Spacer(1,6*mm)]
    for head,rows in sections:
        if head: story.append(Paragraph(esc(head),caption))
        if not rows: continue
        table_data=[[Paragraph(esc(v),cell) for v in row] for row in rows]
        longest=max(len(r) for r in rows)
        available=page[0]-30*mm
        widths=[available/longest]*longest
        if longest==2: widths=[available*0.60,available*0.40]
        table=LongTable(table_data,colWidths=widths,repeatRows=1 if longest>=3 else 0,hAlign="LEFT")
        rules=[("VALIGN",(0,0),(-1,-1),"MIDDLE"),
               ("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6),
               ("LEFTPADDING",(0,0),(-1,-1),7),("RIGHTPADDING",(0,0),(-1,-1),7),
               ("LINEBELOW",(0,0),(-1,0),0.6,colors.HexColor("#DCE5EF"))]
        if longest>=3:
            rules += [("BACKGROUND",(0,0),(-1,0),colors.HexColor("#EAF2FC")),
                      ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F8FAFC")])]
        table.setStyle(TableStyle(rules))
        story.append(table)
    if footer: story.append(Paragraph(esc(footer),foot))
    def draw_page(canvas,document):
        canvas.setFont(font,7.5)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawRightString(page[0]-15*mm,9*mm,f"Trang {document.page}")
    doc.build(story,onFirstPage=draw_page,onLaterPages=draw_page)
    return b.getvalue()


def invoice_pdf(inv: dict[str, Any]) -> bytes:
    s=inv["snapshot"]
    items=[
        ["Khoản thu","Diễn giải","Thành tiền"],
        ["Tiền phòng",period_label(inv["period"]),money(s["rent_amount"])],
        ["Điện",f'{s["e_start"]} → {s["e_end"]} = {s["electric_use"]} kWh × {money(s["e_rate"])}', money(s["electric_amount"])],
        ["Nước",f'{s["w_start"]} → {s["w_end"]} = {s["water_use"]} m³ × {money(s["w_rate"])}',money(s["water_amount"])],
        ["Phụ phí", "Internet / rác / dịch vụ theo cấu hình",money(s["fee_amount"])],
        ["Giảm trừ","Giảm giá trong kỳ", "-"+money(s["discount_amount"])],
        ["TỔNG PHÁT SINH KỲ NÀY","",money(inv["amount"])],
    ]
    bank=" • ".join(str(x) for x in [s.get("bank_name",""),s.get("bank_holder",""),s.get("bank_account","")] if x)
    sections=[("Chi tiết tiền phòng",items),
              ("Công nợ tại thời điểm phát hành",[
                  ["Còn nợ các kỳ trước",money(inv["prior_balance"])],
                  ["Tiền trả trước phân bổ khi phát hành", "-"+money(inv["credit_used"])],
                  ["TỔNG CẦN THANH TOÁN KHI PHÁT HÀNH",money(max(0,int(inv["amount"])+int(inv["prior_balance"])-int(inv["credit_used"])))],
              ])]
    if bank: sections.append(("Thông tin chuyển khoản",[["Ngân hàng / người nhận / số tài khoản",bank],
                 ["Nội dung đề nghị",f'TIEN TRO {s["room_code"]} {inv["period"]}']]))
    return _pdf_doc("THÔNG BÁO TIỀN PHÒNG",f'{s["property_name"]} — {period_label(inv["period"])}\n'
                    f'Phòng: {s["room_code"]} | Khách: {s["tenant"]}\n'
                    f'Mã: TV-{inv["id"]:06d} | Ngày lập: {inv["issued_at"][:10]} | Hạn: {inv["due_at"]}',
                    sections,"Đây là thông báo thu tiền, không phải hóa đơn điện tử thuế.")


def receipt_pdf(payment_id: int, db_path: str=DB_PATH) -> bytes:
    with db_transaction(db_path) as c:
        p=c.execute("SELECT p.*,r.code,r.tenant FROM payments p JOIN rooms r ON r.id=p.room_id WHERE p.id=?",(payment_id,)).fetchone()
        if not p: raise ValueError("Không tìm thấy khoản thanh toán")
        if c.execute("SELECT 1 FROM payment_reversals WHERE payment_id=?",(payment_id,)).fetchone():
            raise ValueError("Giao dịch đã bị đảo; không thể phát hành biên nhận hiệu lực")
        lines=c.execute("""SELECT i.period,a.amount FROM allocations a JOIN invoices i ON i.id=a.invoice_id
                WHERE a.payment_id=? ORDER BY i.period""",(payment_id,)).fetchall()
        cfg={r["key"]:r["value"] for r in c.execute("SELECT key,value FROM settings")}
    allocated=sum(int(l["amount"]) for l in lines)
    rows=[["Khoản","Giá trị"],["Phòng",p["code"]],["Khách",p["tenant"]],
          ["Ngày nhận",p["paid_at"]],["Phương thức",p["method"]],
          ["SỐ TIỀN ĐÃ NHẬN",money(p["amount"])],
          ["Đã phân bổ cho chứng từ",money(allocated)],
          ["Tiền trả trước chưa phân bổ",money(int(p["amount"])-allocated)]]
    payments=[["Kỳ","Số tiền phân bổ"]]+[[period_label(l["period"]),money(l["amount"])] for l in lines]
    return _pdf_doc("BIÊN NHẬN THANH TOÁN",f'{cfg["property_name"]} | Mã biên nhận TV-R{payment_id:06d}',
                    [("Thông tin thực nhận",rows),("Phân bổ thanh toán",payments)],
                    "Biên nhận xác nhận số tiền thực tế đã nhận, không mặc nhiên xác nhận thanh toán toàn bộ công nợ.")


def report_dataframe_rows(period: str, db_path: str=DB_PATH):
    all_inv=invoice_rows(period,db_path)
    output=[]
    for i in all_inv:
        s=i["snapshot"]
        output.append({"Phòng":s["room_code"],"Người thuê":s["tenant"],"Kỳ":i["period"],
                       "Tiền phòng":s["rent_amount"],"Tiền điện":s["electric_amount"],
                       "Tiền nước":s["water_amount"],"Phụ phí":s["fee_amount"],
                       "Giảm trừ":s["discount_amount"],"Phát sinh":i["amount"],
                       "Đã thu phân bổ":i["paid"],"Còn nợ kỳ này":i["outstanding"],
                       "Nợ cũ khi chốt":i["prior_balance"],"Hạn":i["due_at"],
                       "Trạng thái":"Đã đóng đủ" if i["outstanding"]==0 else ("Đóng một phần" if i["paid"] else "Chưa đóng")})
    return output


def export_csv(period: str, db_path: str=DB_PATH) -> bytes:
    rows=report_dataframe_rows(period,db_path)
    cols=["Phòng","Người thuê","Kỳ","Tiền phòng","Tiền điện","Tiền nước","Phụ phí", "Giảm trừ","Phát sinh","Đã thu phân bổ","Còn nợ kỳ này","Nợ cũ khi chốt","Hạn","Trạng thái"]
    output=io.StringIO();writer=csv.DictWriter(output,fieldnames=cols)
    writer.writeheader();writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig")


def export_xlsx(period: str, db_path: str=DB_PATH) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins
    book=Workbook();book.remove(book.active)
    metrics=report_metrics(period,db_path)
    rows=report_dataframe_rows(period,db_path)
    rooms=rooms_list(db_path)
    invs=invoice_rows(db_path=db_path)
    sheets={
        "TONG_QUAN": [["CHỈ TIÊU", "GIÁ TRỊ"],
           ["Kỳ", period_label(period)], ["Tổng phòng",metrics["rooms"]],
           ["Đang thuê",metrics["occupied"]],["Đã phát hành kỳ này",metrics["issued"]],
           ["Đã phân bổ kỳ này",metrics["paid_on_invoices"]],
           ["Thực thu trong tháng",metrics["cash_in_period"]],
           ["Công nợ tất cả kỳ",metrics["outstanding_all"]]],
        "THU_TIEN": ([list(rows[0])] + [list(row.values()) for row in rows]) if rows else [["Phòng","Người thuê","Kỳ","Phát sinh","Đã thu","Còn nợ"]],
        "CONG_NO": [["Phòng","Khách thuê","Kỳ","Còn nợ","Hạn"]] +
           [[i["snapshot"]["room_code"],i["snapshot"]["tenant"],i["period"],i["outstanding"],i["due_at"]]
            for i in invs if i["outstanding"]>0],
        "DIEN_NUOC": [["Phòng","Kỳ","Điện đầu","Điện cuối","kWh","Tiền điện","Nước đầu","Nước cuối","m³","Tiền nước"]] +
           [[i["snapshot"]["room_code"],i["period"],i["snapshot"]["e_start"],i["snapshot"]["e_end"],
             float(i["snapshot"]["electric_use"]),i["snapshot"]["electric_amount"],i["snapshot"]["w_start"],
             i["snapshot"]["w_end"],float(i["snapshot"]["water_use"]),i["snapshot"]["water_amount"]]
            for i in invs if i["period"]==period],
        "DANH_SACH_PHONG": [["Mã phòng","Nhóm","Giá thuê","Trạng thái","Người thuê","Email","Ngày vào","Ngày ra"]] +
             [[r["code"],r["category"],r["rent"],r["status"],r["tenant"],r["email"],r["start_date"],r["end_date"]] for r in rooms]
    }
    for name, values in sheets.items():
        ws=book.create_sheet(name)
        for vals in values: ws.append(vals)
        ws.freeze_panes="A2";ws.sheet_view.showGridLines=False
        ws.auto_filter.ref=ws.dimensions if len(values)>1 else None
        ws.row_dimensions[1].height=27
        for cell in ws[1]:
            cell.fill=PatternFill("solid",fgColor="1769C2")
            cell.font=Font(name="Arial",size=10,color="FFFFFF",bold=True)
            cell.alignment=Alignment(vertical="center",wrap_text=True)
        for col in ws.columns:
            letter=get_column_letter(col[0].column)
            longest=max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[letter].width=max(14,min(35,longest+3))
            for cell in list(col)[1:]:
                cell.font=Font(name="Arial",size=10,color="17324D")
                cell.alignment=Alignment(vertical="center")
                if isinstance(cell.value,(int,float)):
                    cell.number_format='#,##0;[Red](#,##0);–'
        ws.page_margins=PageMargins(left=0.25,right=0.25,top=0.5,bottom=0.5,header=0.2,footer=0.2)
    output=io.BytesIO();book.save(output);return output.getvalue()


def report_pdf(period: str, db_path: str=DB_PATH) -> bytes:
    rows=report_dataframe_rows(period,db_path);m=report_metrics(period,db_path)
    detail=[["Phòng","Khách thuê","Phát sinh","Đã đóng","Còn nợ"]]
    for row in rows: detail.append([row["Phòng"],row["Người thuê"],money(row["Phát sinh"]),money(row["Đã thu phân bổ"]),money(row["Còn nợ kỳ này"])])
    detail.append(["TỔNG","",money(m["issued"]),money(m["paid_on_invoices"]),money(m["issued"]-m["paid_on_invoices"])])
    sections=[("Tổng quan",[["Tổng số phòng",str(m["rooms"])],["Phòng đang thuê",str(m["occupied"])],
               ["Phát sinh kỳ này",money(m["issued"])],["Tiền thu trong tháng",money(m["cash_in_period"])],
               ["Công nợ tất cả kỳ",money(m["outstanding_all"])]]),
              ("Bảng thu tiền",detail)]
    return _pdf_doc("BÁO CÁO QUẢN LÝ TIỀN TRỌ",period_label(period),sections,
                    "Số liệu tổng hợp từ chứng từ đã chốt và giao dịch thanh toán hợp lệ.",True)


def backup_json(db_path: str=DB_PATH) -> bytes:
    with db_transaction(db_path) as c:
        data={name:[dict(x) for x in c.execute(f"SELECT * FROM {name}")] for name in BACKUP_TABLES}
    return json.dumps({"product":"TRO_VIET_MANAGER","version":APP_VERSION,"saved_at":now_iso(),
                       "tables":data},ensure_ascii=False,indent=2).encode("utf-8")


def restore_json(blob: bytes, db_path: str=DB_PATH):
    if len(blob)>20_000_000: raise ValueError("Bản sao lưu vượt giới hạn 20 MB")
    source=json.loads(blob.decode("utf-8-sig"))
    if source.get("product")!="TRO_VIET_MANAGER" or not isinstance(source.get("tables"),dict):
        raise ValueError("Không đúng định dạng sao lưu Trọ Việt Manager")
    tables=source["tables"]
    if set(tables)!=set(BACKUP_TABLES) or not all(isinstance(tables[t],list) for t in BACKUP_TABLES):
        raise ValueError("Thiếu bảng sao lưu bắt buộc")
    with db_transaction(db_path) as c:
        valid={name:{row[1] for row in c.execute(f"PRAGMA table_info({name})")} for name in BACKUP_TABLES}
        for name,items in tables.items():
            for row in items:
                if not isinstance(row,dict) or set(row)!=valid[name]:
                    raise ValueError(f"Cấu trúc sao lưu không khớp bảng {name}")
        for name in reversed(BACKUP_TABLES): c.execute(f"DELETE FROM {name}")
        for name in BACKUP_TABLES:
            for row in tables[name]:
                cols=list(row)
                c.execute(f"INSERT INTO {name} ("+",".join(cols)+") VALUES ("+",".join("?" for _ in cols)+")",tuple(row.values()))
        audit(c,"restore_json",{"version":source.get("version"),"tables":len(tables)})


def smtp_config(secrets: dict[str,Any]|None=None) -> dict[str,str]:
    values={}
    for name in ["SMTP_HOST","SMTP_PORT","SMTP_USER","SMTP_PASSWORD","SMTP_FROM"]:
        values[name]=str((secrets or {}).get(name) or os.environ.get(name) or "").strip()
    return values


def send_invoice_email(invoice_id: int, smtp: dict[str,str], db_path: str=DB_PATH):
    inv=invoice_by_id(invoice_id,db_path);s=inv["snapshot"]
    recipient=s.get("email","").strip()
    if not email_valid(recipient): raise ValueError("Khách chưa có email hợp lệ tại thời điểm chốt")
    if not all(smtp.get(k) for k in ("SMTP_HOST","SMTP_PORT","SMTP_USER","SMTP_PASSWORD","SMTP_FROM")):
        raise ValueError("Email API/SMTP chưa được cấu hình. Hãy tải PDF và gửi thủ công.")
    sent=query("SELECT id FROM email_events WHERE invoice_id=? AND recipient=? AND status='Đã gửi' LIMIT 1",
               (invoice_id,recipient),db_path)
    if sent: raise ValueError("Email đã được gửi cho địa chỉ này. Không gửi trùng tự động.")
    pdf=invoice_pdf(inv)
    msg=EmailMessage()
    msg["Subject"]=f'[TRỌ VIỆT] Thông báo tiền phòng {s["room_code"]} - {period_label(inv["period"])}'
    msg["From"]=smtp["SMTP_FROM"];msg["To"]=recipient
    msg.set_content(f'Kính gửi {s["tenant"]},\n\n'
                    f'Phòng {s["room_code"]} - {period_label(inv["period"])}.\n'
                    f'Tiền phát sinh kỳ này: {money(inv["amount"])}.\n'
                    f'Công nợ trước kỳ: {money(inv["prior_balance"])}.\n'
                    f'Hạn thanh toán: {inv["due_at"]}.\n'
                    'Vui lòng xem PDF đính kèm và liên hệ chủ trọ nếu có sai lệch.\n\nTrân trọng.')
    msg.add_attachment(pdf,maintype="application",subtype="pdf",filename=f'Thong_bao_{s["room_code"]}_{inv["period"]}.pdf')
    status="Đã gửi";message="SMTP đã chấp nhận gửi; chưa xác nhận hộp thư nhận"
    try:
        port=int(smtp["SMTP_PORT"])
        if port==465:
            with smtplib.SMTP_SSL(smtp["SMTP_HOST"],port,timeout=20,context=ssl.create_default_context()) as server:
                server.login(smtp["SMTP_USER"],smtp["SMTP_PASSWORD"]);server.send_message(msg)
        else:
            with smtplib.SMTP(smtp["SMTP_HOST"],port,timeout=20) as server:
                server.starttls(context=ssl.create_default_context())
                server.login(smtp["SMTP_USER"],smtp["SMTP_PASSWORD"]);server.send_message(msg)
    except Exception as e:
        status="Lỗi gửi";message=f"{type(e).__name__}: {str(e)[:170]}"
    with db_transaction(db_path) as c:
        c.execute("INSERT INTO email_events(invoice_id,recipient,status,message,created_at) VALUES(?,?,?,?,?)",
                  (invoice_id,recipient,status,message,now_iso()))
        audit(c,"email_attempt",{"invoice_id":invoice_id,"status":status,"recipient":recipient})
    if status!="Đã gửi": raise RuntimeError(message)
    return message

# --------------------------------------------------------------
# Streamlit interface: six menu items only
# --------------------------------------------------------------

def _secrets():
    try:
        import streamlit as st
        return dict(st.secrets)
    except Exception:
        return {}


def _guard(st):
    credentials=_secrets()
    password=str(credentials.get("APP_PASSWORD") or os.environ.get("APP_PASSWORD") or "")
    if not password:
        st.warning("**Chế độ chưa bảo vệ:** Chưa thiết lập `APP_PASSWORD`. Không nhập dữ liệu cá nhân thật và KHÔNG công khai ứng dụng này trước khi đặt mật khẩu. SQLite trên máy chủ cloud có thể mất khi khởi động lại.")
        return
    if st.session_state.get("authenticated"): return
    st.title("🔐 Trọ Việt Manager")
    st.caption("Vui lòng nhập mật khẩu quản lý được cấu hình trong Streamlit Secrets.")
    with st.form("login"):
        pw=st.text_input("Mật khẩu",type="password")
        submit=st.form_submit_button("Đăng nhập",type="primary")
    if submit:
        if hmac.compare_digest(pw,password):
            st.session_state.authenticated=True
            st.rerun()
        else: st.error("Sai mật khẩu")
    st.stop()


def _styling(st):
    st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Be+Vietnam+Pro:wght@400;500;600;700&display=swap');
    html,body,[class*="css"],[data-testid="stAppViewContainer"]{font-family:'Be Vietnam Pro',Arial,sans-serif!important}
    [data-testid="stAppViewContainer"]{background:#F5F7FA;color:#17324D}
    [data-testid="stSidebar"]{background:#FFFFFF;border-right:1px solid #E3EAF2}
    h1,h2,h3{font-family:'Be Vietnam Pro',Arial,sans-serif!important;color:#17324D;font-weight:700}
    h1{font-size:1.72rem!important}h2{font-size:1.28rem!important}h3{font-size:1.1rem!important}
    [data-testid="stMetric"]{background:white;border:1px solid #E4ECF5;border-radius:13px;padding:16px}
    [data-testid="stMetricValue"]{font-size:1.34rem!important;color:#1769C2}
    div.stButton>button[kind="primary"],div.stDownloadButton>button[kind="primary"]{background:#1769C2;border:0}
    div.stButton>button,div.stDownloadButton>button{border-radius:9px}
    div[data-testid="stDataFrame"]{background:white;border-radius:12px}
    .smallnote{color:#64748b;font-size:0.82rem}
    </style>""",unsafe_allow_html=True)


def _room_title(room):
    return f'{room["code"]} — {room["tenant"] or "Chưa có khách"}'


def _invoice_title(inv):
    return f'TV-{inv["id"]:06d} | {inv["snapshot"]["room_code"]} | {period_label(inv["period"])}'


def _file_dl(st,label,data,name,mime,button_type="secondary",key=None):
    st.download_button(label,data=data,file_name=name,mime=mime,type=button_type,key=key,use_container_width=True)


def screen_dashboard(st,period,db_path):
    metrics=report_metrics(period,db_path)
    st.subheader("Tổng quan · "+period_label(period))
    cols=st.columns(4)
    for col,title,value in zip(cols,["Tổng phòng", "Đang thuê", "Phải thu kỳ này", "Đã thu kỳ này"],
                               [metrics["rooms"],metrics["occupied"],money(metrics["issued"]),money(metrics["paid_on_invoices"])]):
        col.metric(title,value)
    c1,c2,c3,c4=st.columns(4)
    c1.metric("Công nợ mọi kỳ",money(metrics["outstanding_all"]))
    c2.metric("Đã đóng đủ",metrics["paid_full"])
    c3.metric("Đóng một phần",metrics["paid_part"])
    c4.metric("Chưa đóng kỳ này",metrics["unpaid"])
    st.caption("Phải thu = tiền phát sinh theo chứng từ đã chốt; thực thu = tiền đã ghi nhận. Công nợ mọi kỳ không cộng trùng vào doanh thu.")
    st.divider()
    left,right=st.columns([3,2])
    with left:
        st.markdown("#### Tiền phát sinh và đã thu theo kỳ")
        invoices=invoice_rows(db_path=db_path)
        grouped={}
        for r in invoices:
            o=grouped.setdefault(r["period"],{"Phát sinh":0,"Đã phân bổ":0})
            o["Phát sinh"]+=int(r["amount"]);o["Đã phân bổ"]+=int(r["paid"])
        if grouped:
            import pandas as pd
            df=pd.DataFrame([{"Kỳ":k,**v} for k,v in sorted(grouped.items())]).set_index("Kỳ")
            st.bar_chart(df)
        else: st.info("Chưa có thông báo tiền phòng. Bắt đầu bằng cách nhập công tơ và chốt tiền.")
    with right:
        st.markdown("#### Việc cần xử lý")
        rooms=[r for r in rooms_list(db_path) if r["status"]=="Đang thuê"]
        entered={r["room_id"] for r in query("SELECT room_id FROM readings WHERE period=?",(period,),db_path)}
        issued={r["room_id"] for r in invoice_rows(period,db_path)}
        st.write(f"• Chưa nhập điện/nước: **{sum(r['id'] not in entered for r in rooms)} phòng**")
        st.write(f"• Chưa chốt tiền: **{sum(r['id'] not in issued for r in rooms)} phòng**")
        st.write(f"• Phòng còn nợ trong kỳ: **{sum(i['outstanding']>0 for i in invoice_rows(period,db_path))} phòng**")
        events=query("""SELECT COUNT(*) AS total FROM email_events WHERE status='Lỗi gửi'""",db_path=db_path)
        st.write(f"• Lượt gửi email lỗi: **{events[0]['total']}**")
        if st.button("Mở bảng thu tiền",type="primary"):
            st.session_state["nav_target"]="Thu tiền";st.rerun()


def screen_rooms(st,period,db_path):
    st.subheader("Phòng trọ · Khách thuê · Hợp đồng")
    rooms=rooms_list(db_path)
    a,b=st.columns([2,1])
    with a:
        needle=st.text_input("Tìm mã phòng hoặc người thuê",placeholder="Ví dụ: P203")
    with b:
        status_filter=st.selectbox("Lọc trạng thái",["Tất cả","Còn trống","Đang thuê","Đã giữ chỗ","Đang bảo trì","Ngừng khai thác"])
    filtered=[r for r in rooms if (status_filter=="Tất cả" or r["status"]==status_filter) and
              (needle.lower() in (r["code"]+" "+r["tenant"]).lower())]
    import pandas as pd
    st.dataframe(pd.DataFrame([{"Phòng":r["code"],"Loại":r["category"],"Giá thuê":r["rent"],
            "Trạng thái":r["status"],"Khách":r["tenant"],"Email":r["email"]} for r in filtered]),
            hide_index=True,use_container_width=True)
    st.divider()
    mode=st.radio("Thao tác",["Sửa thông tin phòng","Thêm phòng"],horizontal=True)
    existing=(mode=="Sửa thông tin phòng")
    if existing and not rooms: st.info("Chưa có phòng");return
    selected=st.selectbox("Chọn phòng",rooms,format_func=_room_title) if existing else None
    r=selected or {"id":None,"code":"","category":"Tiêu chuẩn","rent":2_000_000,"fee":0,
                   "discount":0,"status":"Còn trống","tenant":"","email":"","phone":"",
                   "start_date":"","end_date":"","deposit":0,"building":"","floor":"",
                   "area":"","notes":"","initial_e":"","initial_w":""}
    with st.form("room_edit_form"):
        c1,c2,c3=st.columns(3)
        with c1:
            code=st.text_input("Mã phòng",r["code"])
            category=st.selectbox("Loại phòng",list(CATEGORIES),index=list(CATEGORIES).index(r["category"]) if r["category"] in CATEGORIES else 0)
            rent=st.number_input("Giá thuê (đ/tháng)",min_value=0,value=int(r["rent"]),step=100000)
            fee=st.number_input("Phụ phí định kỳ (đ)",min_value=0,value=int(r["fee"]),step=10000)
            discount=st.number_input("Giảm trừ định kỳ (đ)",min_value=0,value=int(r["discount"]),step=10000)
        with c2:
            status_values=["Còn trống","Đang thuê","Đã giữ chỗ","Đang bảo trì","Ngừng khai thác"]
            status=st.selectbox("Trạng thái",status_values,index=status_values.index(r["status"]) if r["status"] in status_values else 0)
            tenant=st.text_input("Người thuê",r["tenant"])
            email=st.text_input("Email nhận PDF",r["email"])
            phone=st.text_input("Số điện thoại",r["phone"])
            deposit=st.number_input("Tiền đặt cọc (đ)",min_value=0,value=int(r["deposit"]),step=100000)
        with c3:
            start=st.text_input("Ngày bắt đầu (dd/MM/yyyy)",display_date(r["start_date"]),
                                placeholder="Ví dụ: 15/01/2026")
            end=st.text_input("Ngày kết thúc (dd/MM/yyyy, có thể để trống)",display_date(r["end_date"]),
                              placeholder="Ví dụ: 31/12/2026")
            building=st.text_input("Dãy/tòa",r["building"])
            floor=st.text_input("Tầng",r["floor"])
            area=st.text_input("Diện tích (m²)",r["area"])
            e0=st.text_input("Số điện đầu tiên đã xác minh",r["initial_e"])
            w0=st.text_input("Số nước đầu tiên đã xác minh",r["initial_w"])
        notes=st.text_area("Ghi chú",r["notes"],height=70)
        save=st.form_submit_button("Lưu thông tin phòng",type="primary")
    if save:
        try:
            values=dict(id=r["id"],code=code.strip(),category=category,rent=rent,fee=fee,discount=discount,
                        status=status,tenant=tenant.strip(),email=email.strip(),phone=phone.strip(),
                        start_date=start,end_date=end,deposit=deposit,building=building,floor=floor,area=area,
                        notes=notes,initial_e=e0.strip(),initial_w=w0.strip())
            save_room(values,db_path)
            st.success("Đã lưu và tạo lịch sử hợp đồng khi có thay đổi.");st.rerun()
        except Exception as e: st.error(str(e))
    if existing:
        contracts=query("SELECT * FROM contracts WHERE room_id=? ORDER BY id DESC",(r["id"],),db_path)
        with st.expander("Lịch sử khách thuê / hợp đồng"):
            if contracts: st.dataframe(pd.DataFrame(contracts),hide_index=True)
            else: st.info("Chưa có phiên bản hợp đồng nào được lưu")
        with st.expander("Ngừng khai thác / xóa mềm phòng"):
            if st.checkbox("Tôi xác nhận muốn ẩn phòng này khỏi danh sách quản lý",key=f"del_{r['id']}"):
                if st.button("Ẩn phòng",key=f"remove_{r['id']}"):
                    delete_room(r["id"],db_path);st.rerun()


def screen_meters(st,period,db_path):
    st.subheader("Điện & nước · "+period_label(period))
    rooms=[r for r in rooms_list(db_path) if r["status"]=="Đang thuê"]
    if not rooms:
        st.info("Chưa có phòng đang thuê. Hãy cập nhật Phòng trọ và số điện/nước đầu kỳ.");return
    row=st.selectbox("Chọn phòng để ghi công tơ",rooms,format_func=_room_title)
    stored=query("SELECT * FROM readings WHERE room_id=? AND period=?",(row["id"],period),db_path)
    read=stored[0] if stored else None
    try: e_start=read["e_start"] if read else meter_start(row["id"],period,"e",db_path)
    except Exception:e_start=""
    try: w_start=read["w_start"] if read else meter_start(row["id"],period,"w",db_path)
    except Exception:w_start=""
    if not e_start or not w_start:
        st.warning("Phòng chưa có số đầu kỳ từ tháng đã chốt. Hãy nhập đúng số cũ thực tế; có thể lưu số khởi tạo trong Phòng trọ.")
    with st.form("meter_form"):
        c1,c2=st.columns(2)
        with c1:
            old_e=st.number_input("Điện đầu kỳ (kWh)",value=float(e_start or 0),min_value=0.0,step=1.0,format="%.3f")
            new_e=st.number_input("Điện cuối kỳ (kWh)",value=float(read["e_end"] if read else e_start or 0),min_value=0.0,step=1.0,format="%.3f")
        with c2:
            old_w=st.number_input("Nước đầu kỳ (m³)",value=float(w_start or 0),min_value=0.0,step=0.1,format="%.3f")
            new_w=st.number_input("Nước cuối kỳ (m³)",value=float(read["w_end"] if read else w_start or 0),min_value=0.0,step=0.1,format="%.3f")
        if st.form_submit_button("Lưu chỉ số công tơ",type="primary"):
            try:
                save_reading(row["id"],period,old_e,new_e,old_w,new_w,db_path)
                st.success("Đã lưu số công tơ");st.rerun()
            except Exception as e: st.error(str(e))
    cfg=settings_get(db_path)
    readings=query("""SELECT r.code,r.tenant,m.room_id,m.e_start,m.e_end,m.w_start,m.w_end
                      FROM readings m JOIN rooms r ON r.id=m.room_id WHERE m.period=? ORDER BY r.code""",(period,),db_path)
    if readings:
        st.markdown("#### Bảng công tơ đã nhập")
        import pandas as pd
        data=[]
        for r in readings:
            det=calculate_bill(0,r["e_start"],r["e_end"],r["w_start"],r["w_end"],int(cfg["electric_rate"]),int(cfg["water_rate"]),0)
            data.append({"Phòng":r["code"],"Khách thuê":r["tenant"],"Điện đầu":r["e_start"],"Điện cuối":r["e_end"],
                         "kWh":det["electric_use"],"Tiền điện":det["electric_amount"],"Nước đầu":r["w_start"],
                         "Nước cuối":r["w_end"],"m³":det["water_use"],"Tiền nước":det["water_amount"]})
        st.dataframe(pd.DataFrame(data),hide_index=True,use_container_width=True)
        _file_dl(st,"Tải CSV công tơ",pd.DataFrame(data).to_csv(index=False).encode("utf-8-sig"),f"Cong_to_{period}.csv","text/csv",key="readings_download")
    with st.expander("Nhập nhiều phòng nhanh từ CSV"):
        st.caption("Mỗi dòng: Mã phòng;Điện đầu;Điện cuối;Nước đầu;Nước cuối. Phòng đã chốt không bị thay đổi. Hãy kiểm tra số liệu trước khi nhập.")
        raw=st.text_area("Dữ liệu (mỗi phòng một dòng)",placeholder="P001;1250;1335;120;125\nP002;200;220;30;32")
        if st.button("Kiểm tra và nhập nhiều phòng",key="bulk_meter_btn") and raw.strip():
            mapping={r["code"]:r for r in rooms};count=0;errors=[]
            for line in raw.splitlines():
                parts=[x.strip() for x in line.split(";")]
                if len(parts)!=5 or parts[0] not in mapping:
                    errors.append(line+" → Sai định dạng hoặc mã phòng");continue
                try: save_reading(mapping[parts[0]]["id"],period,*parts[1:],db_path);count+=1
                except Exception as e: errors.append(parts[0]+" → "+str(e))
            st.success(f"Đã lưu {count} phòng")
            for error in errors: st.error(error)


def _status(inv):
    return "Đã đóng đủ" if inv["outstanding"]==0 else ("Đã đóng một phần" if inv["paid"] else "Chưa đóng")


def screen_billing(st,period,db_path):
    st.subheader("Thu tiền · "+period_label(period))
    tabs=st.tabs(["Chốt tiền & thông báo PDF", "Xác nhận đã đóng", "Công nợ & lịch sử", "Chuyển tháng"])
    with tabs[0]:
        st.markdown("#### 1. Xác nhận tiền cần thu")
        st.caption("Xác nhận tiền cần thu chỉ phát hành thông báo, KHÔNG ghi nhận khách đã đóng.")
        rooms=[r for r in rooms_list(db_path) if r["status"]=="Đang thuê"]
        issued={i["room_id"] for i in invoice_rows(period,db_path)}
        eligible=[r for r in rooms if r["id"] not in issued]
        if eligible:
            target=st.selectbox("Chọn phòng chưa chốt",eligible,format_func=_room_title)
            read=query("SELECT * FROM readings WHERE room_id=? AND period=?",(target["id"],period),db_path)
            if read:
                c=settings_get(db_path)
                try:
                    det=calculate_bill(rent_for_period(target["rent"],period,target["start_date"],target["end_date"]),
                        read[0]["e_start"],read[0]["e_end"],read[0]["w_start"],read[0]["w_end"],
                        int(c["electric_rate"]),int(c["water_rate"]),target["fee"],target["discount"])
                    k1,k2,k3,k4=st.columns(4)
                    k1.metric("Tiền phòng",money(det["rent_amount"]))
                    k2.metric("Tiền điện",money(det["electric_amount"]))
                    k3.metric("Tiền nước",money(det["water_amount"]))
                    k4.metric("Phát sinh",money(det["amount"]))
                except ValueError as e:st.error(str(e))
            else: st.warning("Phòng chưa nhập công tơ kỳ này")
            valid=st.checkbox("Tôi đã kiểm tra tiền phòng, điện, nước và thông tin người thuê",key="confirm_bill")
            if st.button("Xác nhận tiền cần thu · Chốt PDF",type="primary",disabled=not valid):
                try:
                    invoice_id=issue_invoice(target["id"],period,db_path)
                    st.success(f"Đã chốt thông báo TV-{invoice_id:06d}. Có thể tải PDF phía dưới.")
                    cfg=settings_get(db_path)
                    if cfg.get("auto_email")=="1":
                        try: st.info(send_invoice_email(invoice_id,smtp_config(_secrets()),db_path))
                        except Exception as e: st.warning("Đã chốt PDF nhưng email chưa gửi: "+str(e))
                    st.rerun()
                except Exception as e:st.error(str(e))
            with st.expander("Chốt nhiều phòng đủ điều kiện"):
                issue_all=st.checkbox("Chọn tất cả phòng đang thuê chưa chốt",key="issue_all_rooms")
                selected=eligible if issue_all else st.multiselect("Chọn các phòng",eligible,format_func=_room_title,key="bulk_issue_rooms")
                approve=st.checkbox("Xác nhận chốt những phòng đã chọn",key="bulk_issue_confirm")
                if st.button("Chốt các phòng đã chọn",disabled=not (selected and approve)):
                    done=[];failed=[]
                    for rr in selected:
                        try:
                            inv_id=issue_invoice(rr["id"],period,db_path);done.append(rr["code"])
                            if settings_get(db_path).get("auto_email")=="1":
                                try:send_invoice_email(inv_id,smtp_config(_secrets()),db_path)
                                except Exception as e: failed.append(rr["code"]+" (email: "+str(e)+")")
                        except Exception as e:failed.append(rr["code"]+": "+str(e))
                    st.success("Đã chốt: "+", ".join(done) if done else "Không có phòng nào được chốt")
                    for f in failed:st.warning(f)
        else: st.info("Không có phòng đang thuê chưa chốt tiền trong kỳ này.")
        st.divider()
        current=invoice_rows(period,db_path)
        st.markdown("#### 2. Thông báo đã phát hành")
        if current:
            import pandas as pd
            st.dataframe(pd.DataFrame([{"Mã":f'TV-{i["id"]:06d}',"Phòng":i["snapshot"]["room_code"],
               "Khách":i["snapshot"]["tenant"],"Phát sinh":i["amount"],"Đã thanh toán":i["paid"],
               "Còn nợ":i["outstanding"],"Trạng thái":_status(i)} for i in current]),
               hide_index=True,use_container_width=True)
            selected_inv=st.selectbox("Thông báo cần tải / gửi email",current,format_func=_invoice_title)
            z1,z2=st.columns(2)
            try:
                with z1: _file_dl(st,"Tải PDF thông báo",invoice_pdf(selected_inv),
                                f'Thong_bao_{selected_inv["snapshot"]["room_code"]}_{period}.pdf',"application/pdf",key="invoice_pdf")
            except Exception as e: st.error("Lỗi xuất PDF: "+str(e))
            with z2:
                if st.button("Gửi PDF đến email khách",key="send_one_email"):
                    try: st.success(send_invoice_email(selected_inv["id"],smtp_config(_secrets()),db_path))
                    except Exception as e: st.error(str(e))
            if not all(smtp_config(_secrets()).values()):
                address=selected_inv["snapshot"].get("email","")
                st.info("SMTP chưa cấu hình: tải PDF rồi đính kèm thủ công vào email. Nút gửi trực tiếp chỉ hoạt động khi có SMTP thật.")
                st.code(f'Đến: {address or "[EMAIL KHÁCH]"}\nChủ đề: Thông báo tiền phòng {selected_inv["snapshot"]["room_code"]} - {period_label(period)}',language="text")
            with st.expander("Gửi email hàng loạt (chỉ khi có SMTP)"):
                possible=[i for i in current if email_valid(i["snapshot"].get("email",""))]
                email_all=st.checkbox("Chọn tất cả phòng có email hợp lệ",key="email_all_rooms")
                selected_invs=possible if email_all else st.multiselect("Chọn thông báo",possible,format_func=_invoice_title)
                bulk_confirm=st.checkbox("Tôi xác nhận gửi PDF riêng đến đúng email của từng khách",key="bulk_email_ok")
                if st.button("Gửi những email đã chọn",disabled=not(selected_invs and bulk_confirm)):
                    done=[];failed=[]
                    for inv in selected_invs:
                        try:send_invoice_email(inv["id"],smtp_config(_secrets()),db_path);done.append(inv["snapshot"]["room_code"])
                        except Exception as e:failed.append(inv["snapshot"]["room_code"]+": "+str(e))
                    st.success(f"SMTP chấp nhận {len(done)} email")
                    for error in failed:st.warning(error)
        else:st.info("Chưa có thông báo thu tiền đã chốt.")
    with tabs[1]:
        st.markdown("#### Xác nhận đã nhận tiền thực tế")
        all_rooms=rooms_list(db_path)
        debts=[r for r in all_rooms if debt_room(r["id"],db_path)>0]
        if debts:
            room=st.selectbox("Chọn phòng để nhận tiền",debts,format_func=_room_title,key="payment_room")
            remain=debt_room(room["id"],db_path)
            m_col,c_col=st.columns(2)
            m_col.metric("Còn phải thu của phòng (tất cả kỳ)",money(remain))
            c_col.metric("Tiền trả trước chưa phân bổ",money(credit_room(room["id"],db_path)))
            with st.form("payment_form"):
                amount=st.number_input("Số tiền thực nhận (đ)",min_value=1,value=int(remain),step=10000)
                paid_at=st.date_input("Ngày nhận tiền",value=vn_today())
                method=st.selectbox("Phương thức",["Chuyển khoản","Tiền mặt","Khác"])
                note=st.text_input("Ghi chú")
                if st.form_submit_button("Xác nhận đã đóng",type="primary"):
                    try:
                        key=st.session_state.setdefault("payment_token",uuid4().hex)
                        pay_id=record_payment(room["id"],amount,paid_at.isoformat(),method,note,key,db_path)
                        st.session_state["payment_token"]=uuid4().hex
                        st.success(f"Đã ghi nhận giao dịch #{pay_id}. Có thể tải biên nhận ở phần Công nợ & lịch sử.")
                    except Exception as e:st.error(str(e))
            with st.expander("Xác nhận đóng đủ nhiều phòng"):
                selected=st.multiselect("Chọn phòng đã thanh toán đủ",debts,format_func=_room_title,key="bulk_paid_rooms")
                if selected:
                    st.write("Số tiền sẽ ghi nhận:",money(sum(debt_room(r["id"],db_path) for r in selected)))
                bulk_date=st.date_input("Ngày thu",vn_today(),key="bulk_paid_date")
                bulk_method=st.selectbox("Phương thức chung",["Chuyển khoản","Tiền mặt","Khác"],key="bulk_paid_method")
                approved=st.checkbox("Đã kiểm tra chứng từ thanh toán của các phòng",key="bulk_paid_approve")
                if st.button("Xác nhận đã đóng các phòng đã chọn",disabled=not(selected and approved)):
                    completed=[];failed=[]
                    for r in selected:
                        try:
                            amount=debt_room(r["id"],db_path)
                            if amount>0:
                                token=st.session_state.setdefault("bulk_payment_token_"+str(r["id"]),uuid4().hex)
                                record_payment(r["id"],amount,bulk_date.isoformat(),bulk_method,"Xác nhận hàng loạt",token,db_path)
                                st.session_state["bulk_payment_token_"+str(r["id"])]=uuid4().hex
                                completed.append(r["code"])
                        except Exception as e:failed.append(r["code"]+": "+str(e))
                    st.success("Đã ghi nhận: "+", ".join(completed))
                    for f in failed: st.warning(f)
        else:st.info("Chưa có công nợ. Cần chốt thông báo tiền phòng trước khi xác nhận đã đóng.")
    with tabs[2]:
        st.markdown("#### Công nợ & lịch sử giao dịch")
        current_all=[i for i in invoice_rows(db_path=db_path) if i["outstanding"]>0]
        if current_all:
            import pandas as pd
            st.dataframe(pd.DataFrame([{"Phòng":i["snapshot"]["room_code"],"Kỳ":i["period"],
                  "Khách":i["snapshot"]["tenant"],"Còn nợ":i["outstanding"],"Hạn":i["due_at"]}
                  for i in current_all]),hide_index=True,use_container_width=True)
        else:st.success("Không có chứng từ còn nợ")
        payments=payments_list(db_path)
        if payments:
            st.dataframe(pd.DataFrame([{"Mã":p["id"],"Phòng":p["code"],"Ngày thu":p["paid_at"],
                "Số tiền":p["amount"],"Phương thức":p["method"],
                "Trạng thái":"Đã đảo" if p["reversed_at"] else "Hợp lệ"} for p in payments]),
                use_container_width=True,hide_index=True)
            p=st.selectbox("Giao dịch để in biên nhận / điều chỉnh",payments,format_func=lambda p:f'#{p["id"]} · {p["code"]} · {money(p["amount"])}')
            if not p["reversed_at"]:
                try:_file_dl(st,"Tải PDF biên nhận",receipt_pdf(p["id"],db_path),
                            f'Bien_nhan_{p["id"]}.pdf',"application/pdf",key="receipt_dl")
                except Exception as e:st.error(str(e))
                with st.expander("Hủy xác nhận đã đóng nhầm"):
                    reason=st.text_input("Lý do điều chỉnh",key="rev_reason")
                    check=st.checkbox("Tôi xác nhận cần đảo giao dịch này",key="rev_confirm")
                    if st.button("Đảo giao dịch",disabled=not check):
                        try:reverse_payment(p["id"],reason,db_path);st.success("Đã ghi nhận bút toán đảo");st.rerun()
                        except Exception as e:st.error(str(e))
        else:st.info("Chưa có giao dịch thanh toán.")
    with tabs[3]:
        st.markdown("#### Hoàn tất kỳ và chuyển tháng")
        m=report_metrics(period,db_path)
        x,y,z=st.columns(3)
        x.metric("Hóa đơn đã chốt",m["invoices"])
        y.metric("Đã đóng đủ",m["paid_full"])
        z.metric("Công nợ mọi kỳ",money(m["outstanding_all"]))
        st.info("Có thể chuyển tháng ngay cả khi một số phòng chưa thanh toán. Hóa đơn và giao dịch kỳ cũ vẫn được giữ lại.")
        approved=st.checkbox("Tôi đã kiểm tra dữ liệu và muốn mở kỳ mới",key="next_period_confirm")
        if st.button("Hoàn tất kỳ và chuyển sang tháng tiếp theo",type="primary",disabled=not approved):
            with db_transaction(db_path) as c:audit(c,"open_next_period",{"from":period,"to":next_period(period)})
            st.session_state["requested_period"]=next_period(period)
            st.rerun()


def screen_reports(st,period,db_path):
    st.subheader("Báo cáo & xuất dữ liệu")
    metrics=report_metrics(period,db_path)
    a,b,c=st.columns(3)
    a.metric("Phát sinh kỳ này",money(metrics["issued"]))
    b.metric("Tiền thu trong tháng",money(metrics["cash_in_period"]))
    c.metric("Tổng công nợ",money(metrics["outstanding_all"]))
    rows=report_dataframe_rows(period,db_path)
    if rows:
        import pandas as pd
        st.dataframe(pd.DataFrame(rows),hide_index=True,use_container_width=True)
    else:st.info("Kỳ này chưa có thông báo đã phát hành.")
    st.caption("Xuất CSV/Excel/PDF chỉ lấy dữ liệu chứng từ đã chốt; Excel gồm 5 sheet Tổng quan, Thu tiền, Công nợ, Điện nước và Phòng.")
    d1,d2,d3=st.columns(3)
    with d1:
        _file_dl(st,"Xuất Excel 5 sheet",export_xlsx(period,db_path),f'Bao_cao_tong_hop_{period}.xlsx',
                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",key="xlsx")
    with d2:
        _file_dl(st,"Xuất bảng thu CSV",export_csv(period,db_path),f'Bang_thu_tien_{period}.csv',"text/csv",key="csv")
    with d3:
        try:_file_dl(st,"Xuất báo cáo PDF",report_pdf(period,db_path),f'Bao_cao_{period}.pdf',"application/pdf",key="report_pdf")
        except Exception as e:st.error("Không thể tạo PDF: "+str(e))


def screen_settings(st,period,db_path):
    st.subheader("Cài đặt")
    cfg=settings_get(db_path)
    with st.form("settings_form"):
        c1,c2=st.columns(2)
        with c1:
            name=st.text_input("Tên nhà trọ",cfg["property_name"])
            address=st.text_input("Địa chỉ",cfg["property_address"])
            phone=st.text_input("Điện thoại liên hệ",cfg["property_phone"])
            e_rate=st.number_input("Giá điện (đ/kWh)",0,value=int(cfg["electric_rate"]),step=100)
            w_rate=st.number_input("Giá nước (đ/m³)",0,value=int(cfg["water_rate"]),step=500)
        with c2:
            bank=st.text_input("Tên ngân hàng",cfg["bank_name"])
            holder=st.text_input("Chủ tài khoản",cfg["bank_holder"])
            account=st.text_input("Số tài khoản",cfg["bank_account"])
            due_day=st.number_input("Ngày đến hạn (tháng kế tiếp)",min_value=1,max_value=28,value=int(cfg["due_day"]))
            auto_email=st.checkbox("Tự gửi email khi chốt tiền (cần SMTP)",value=cfg.get("auto_email")=="1")
        if st.form_submit_button("Lưu cài đặt",type="primary"):
            try:settings_save({"property_name":name,"property_address":address,"property_phone":phone,
                  "electric_rate":str(e_rate),"water_rate":str(w_rate),"bank_name":bank,
                  "bank_holder":holder,"bank_account":account,"due_day":str(due_day),
                  "auto_email":"1" if auto_email else "0"},db_path)
            except Exception as e:st.error(str(e))
            else:st.success("Đã lưu cấu hình. Chứng từ cũ không bị thay đổi.");st.rerun()
    st.divider()
    st.markdown("#### Sao lưu và khôi phục")
    _file_dl(st,"Tải bản sao lưu JSON",backup_json(db_path),
             f'Tro_Viet_Backup_{vn_today()}.json',"application/json",key="backup")
    st.warning("Khôi phục ghi đè toàn bộ cơ sở dữ liệu đang dùng. Hãy tải bản sao lưu hiện tại trước khi thực hiện.")
    uploaded=st.file_uploader("Chọn file JSON để khôi phục",type=["json"],key="restore_file")
    confirm=st.checkbox("Tôi hiểu rằng khôi phục sẽ thay thế dữ liệu hiện tại",key="restore_confirm")
    if st.button("Khôi phục từ JSON",disabled=not(uploaded and confirm)):
        try:restore_json(uploaded.read(),db_path);st.success("Khôi phục thành công");st.rerun()
        except Exception as e:st.error(str(e))
    st.divider()
    st.markdown("#### Trạng thái tích hợp")
    smtp=smtp_config(_secrets())
    st.write("**Email SMTP:**", "Đã cấu hình" if all(smtp.values()) else "Chưa kết nối — có thể tải PDF và gửi thủ công")
    st.write("**Supabase / đa tài khoản:** Chưa kết nối")
    st.write("**Google Login:** Chưa kết nối")
    st.write("**Google Sheets:** Chưa kết nối")
    st.write("**SMS:** Chưa kết nối")
    st.info("Google Sheets đăng ký chủ trọ và Supabase cần backend, xác thực và quyền truy cập riêng; Streamlit bản này chưa thực hiện đồng bộ thật.")
    st.caption("Dữ liệu SQLite nằm trên máy chủ đang chạy Streamlit. Streamlit Community Cloud có bộ nhớ tệp không bền vững: hãy sao lưu JSON thường xuyên hoặc dùng database được lưu trữ bền vững trước khi dùng thật.")
    st.caption("Mật khẩu đơn giản APP_PASSWORD bảo vệ một chủ trọ; đây chưa phải hệ thống SaaS đa người dùng có phân quyền theo tổ chức.")


def main():
    import streamlit as st
    st.set_page_config(page_title="Trọ Việt Manager",page_icon="🏠",layout="wide",initial_sidebar_state="expanded")
    _styling(st)
    try:init_db(DB_PATH)
    except Exception as e:
        st.error("Không mở được cơ sở dữ liệu: "+str(e));st.stop()
    _guard(st)
    st.sidebar.markdown("### 🏠 TRỌ VIỆT")
    st.sidebar.caption("Quản lý dễ dàng · Thu tiền chính xác")
    if "nav_target" in st.session_state:
        st.session_state["nav_main"]=st.session_state.pop("nav_target")
    if "nav_main" not in st.session_state:
        st.session_state["nav_main"]="Tổng quan"
    menu=st.sidebar.radio("Điều hướng",MENUS,key="nav_main")
    if "requested_period" in st.session_state:
        req=st.session_state.pop("requested_period")
        y,m=map(int,req.split("-"))
        st.session_state["year_sel"]=y
        st.session_state["month_sel"]=m
    today=vn_today()
    st.session_state.setdefault("year_sel",today.year)
    st.session_state.setdefault("month_sel",today.month)
    y_col,m_col=st.sidebar.columns(2)
    with y_col:
        year=st.selectbox("Năm",list(range(today.year-3,today.year+7)),key="year_sel")
    with m_col:
        month=st.selectbox("Tháng",list(range(1,13)),key="month_sel")
    period=f"{year:04d}-{month:02d}"
    st.sidebar.divider()
    st.sidebar.caption("● Chế độ quản lý một chủ · SQLite cục bộ")
    if os.environ.get("APP_PASSWORD") or _secrets().get("APP_PASSWORD"):
        if st.sidebar.button("Đăng xuất"):
            st.session_state.authenticated=False;st.rerun()
    try:
        if menu=="Tổng quan":screen_dashboard(st,period,DB_PATH)
        elif menu=="Phòng trọ":screen_rooms(st,period,DB_PATH)
        elif menu=="Điện & nước":screen_meters(st,period,DB_PATH)
        elif menu=="Thu tiền":screen_billing(st,period,DB_PATH)
        elif menu=="Báo cáo":screen_reports(st,period,DB_PATH)
        else:screen_settings(st,period,DB_PATH)
    except Exception as exc:
        st.error("Đã có lỗi khi mở chức năng: "+str(exc))
        st.caption("Không làm mới dữ liệu bằng cách xóa trình duyệt. Hãy kiểm tra file JSON sao lưu hoặc nhật ký máy chủ.")


if __name__=="__main__":
    main()
